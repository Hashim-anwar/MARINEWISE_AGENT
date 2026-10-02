"""MarineWise AI - builds downloadable files (PPTX, DOCX, PDF, Quiz PDF). All return bytes."""
import io
import os
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches as DInches
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


def sections(c: dict) -> list:
    """Turn training content into a flat list of (heading, bullets, image, source)."""
    out = [("Learning Objectives", c.get("objectives", []), "", ""),
           ("System Overview", [c.get("overview", "")], "", "")]
    for s in c.get("steps", []):
        out.append((s.get("title", "Step"), s.get("points", []), s.get("image", ""), s.get("source", "")))
    out += [("Safety Precautions", c.get("safety", []), "", ""),
            ("Common Mistakes", c.get("mistakes", []), "", ""),
            ("Practical Exercise", [c.get("exercise", "")], "", ""),
            ("Practice Questions", [f"Q: {p['q']}  A: {p['a']}" for p in c.get("practice", [])], "", ""),
            ("Sources (Document | Page | Drive path)", c.get("sources", []), "", "")]
    return [s for s in out if any(s[1])]


def build_pptx(c: dict) -> bytes:
    """Training presentation: title, then one slide per section (6 bullets max per slide)."""
    prs = Presentation()
    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = c.get("title", "Training")
    t.placeholders[1].text = c.get("subtitle", "")
    for head, points, img, src in sections(c):
        for part in range(0, len(points), 6):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = head
            body = s.placeholders[1]
            has_img = bool(img) and os.path.exists(img) and part == 0
            body.left, body.top, body.height = Inches(0.5), Inches(1.5), Inches(5.2)
            body.width = Inches(4.6 if has_img else 9)
            for n, p in enumerate(points[part:part + 6]):
                para = body.text_frame.paragraphs[0] if n == 0 else body.text_frame.add_paragraph()
                para.text = p
                para.font.size = Pt(15)
            if has_img:
                s.shapes.add_picture(img, Inches(5.3), Inches(1.8), width=Inches(4.2))
                note = s.shapes.add_textbox(Inches(5.3), Inches(5.6), Inches(4.2), Inches(0.4))
                note.text_frame.text = "Illustration from OEM manual"
                note.text_frame.paragraphs[0].font.size = Pt(9)
            if src:
                box = s.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(9), Inches(0.4))
                box.text_frame.text = f"Source: {src}"
                box.text_frame.paragraphs[0].font.size = Pt(10)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_docx(c: dict) -> bytes:
    """Detailed training manual in Word format."""
    d = Document()
    d.add_heading(c.get("title", "Training"), 0)
    d.add_paragraph(c.get("subtitle", ""))
    for head, points, img, src in sections(c):
        d.add_heading(head, 1)
        for p in points:
            d.add_paragraph(p, style="List Bullet")
        if img and os.path.exists(img):
            d.add_picture(img, width=DInches(3.8))
        if src:
            d.add_paragraph(f"Source: {src}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_pdf(c: dict) -> bytes:
    """Printable training package in PDF format."""
    st = getSampleStyleSheet()
    story = [Paragraph(escape(c.get("title", "Training")), st["Title"]),
             Paragraph(escape(c.get("subtitle", "")), st["Normal"]), Spacer(1, 12)]
    for head, points, img, src in sections(c):
        story.append(Paragraph(escape(head), st["Heading2"]))
        for p in points:
            story.append(Paragraph("&bull; " + escape(p), st["Normal"]))
        if img and os.path.exists(img):
            story.append(Image(img, width=8 * cm, height=6 * cm, kind="proportional"))
        if src:
            story.append(Paragraph(f"<i>Source: {escape(src)}</i>", st["Normal"]))
        story.append(Spacer(1, 8))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def build_quiz_pdf(quiz: list, title: str) -> bytes:
    """Quiz PDF: questions, then an ANSWER SHEET page, then the ANSWER KEY (Q|Answer|Explanation|Source)."""
    st = getSampleStyleSheet()
    story = [Paragraph(escape(title), st["Title"]),
             Paragraph("Name: ____________________   Date: ____________", st["Normal"]),
             Paragraph("Write your answers clearly on the ANSWER SHEET page.", st["Normal"]),
             Spacer(1, 10)]
    for i, q in enumerate(quiz, start=1):
        story.append(Paragraph(f"<b>{i}.</b> {escape(q['q'])}", st["Normal"]))
        for letter, opt in zip("ABCD", q.get("options") or []):
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;{letter}. {escape(str(opt))}", st["Normal"]))
        if q["type"] == "True/False":
            story.append(Paragraph("&nbsp;&nbsp;&nbsp;True / False", st["Normal"]))
        story.append(Spacer(1, 8))
    story += [PageBreak(), Paragraph("ANSWER SHEET", st["Title"])]
    rows = [["Q", "Your answer"]] + [[str(i), ""] for i in range(1, len(quiz) + 1)]
    sheet = Table(rows, colWidths=[2 * cm, 13 * cm], rowHeights=[0.8 * cm] + [1.1 * cm] * len(quiz))
    sheet.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.8, colors.black)]))
    story += [sheet, PageBreak(), Paragraph("ANSWER KEY (Trainer copy)", st["Title"])]
    rows = [["Q", "Answer", "Explanation", "Source"]]
    for i, q in enumerate(quiz, start=1):
        ans = str(q["answer"])
        if q["type"] == "MCQ" and q.get("options") and ans[:1].upper() in "ABCD":
            idx = "ABCD".index(ans[:1].upper())
            ans = f"{ans[:1].upper()} - {q['options'][idx]}" if idx < len(q["options"]) else ans
        rows.append([str(i)] + [Paragraph(escape(x), st["Normal"])
                                for x in (ans, q.get("explanation", ""), q.get("source", ""))])
    key = Table(rows, colWidths=[1 * cm, 4 * cm, 7.5 * cm, 4 * cm], repeatRows=1)
    key.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                             ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                             ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(key)
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()
